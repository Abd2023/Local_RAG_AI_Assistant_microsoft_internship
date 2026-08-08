from pathlib import Path

import pymupdf
from docx import Document

from src.document_loaders import (
    _needs_ocr,
    iter_supported_files,
    load_docx_document,
    load_document,
    load_pdf_document,
)


def test_loader_registry_finds_supported_files(tmp_path: Path) -> None:
    (tmp_path / "b.txt").write_text("Text", encoding="utf-8")
    (tmp_path / "a.md").write_text("# Title", encoding="utf-8")
    (tmp_path / "ignored.json").write_text("{}", encoding="utf-8")

    assert [path.name for path in iter_supported_files(tmp_path)] == ["a.md", "b.txt"]
    assert load_document(tmp_path / "a.md")[0].extraction_method == "markdown"
    assert load_document(tmp_path / "b.txt")[0].extraction_method == "text"


def test_docx_loader_reads_paragraph_text(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    document = Document()
    document.add_heading("Handbook", level=1)
    document.add_paragraph("The daily standup starts at 10:00 AM.")
    document.save(path)

    blocks = load_docx_document(path)

    assert len(blocks) == 1
    assert blocks[0].source_name == "handbook.docx"
    assert blocks[0].extraction_method == "docx"
    assert "daily standup" in blocks[0].content


def test_pdf_loader_reads_text_pages(tmp_path: Path) -> None:
    path = tmp_path / "handbook.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "The final demo lasts five minutes and includes sources.")
    document.save(path)
    document.close()

    blocks = load_pdf_document(path)

    assert len(blocks) == 1
    assert blocks[0].page_start == 1
    assert blocks[0].page_end == 1
    assert blocks[0].extraction_method == "pdf_text"
    assert "final demo" in blocks[0].content


def test_ocr_decision_uses_minimum_text_threshold() -> None:
    assert _needs_ocr("")
    assert _needs_ocr("short")
    assert not _needs_ocr("This page has enough extracted text to avoid OCR fallback.")
