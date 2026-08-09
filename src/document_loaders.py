"""Document loader registry for supported local source files."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from src import config

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}


class DocumentLoadError(RuntimeError):
    """Raised when a source document cannot be loaded."""


class OCRUnavailableError(DocumentLoadError):
    """Raised when OCR is needed but Tesseract is not available/configured."""


@dataclass(frozen=True)
class LoadedTextBlock:
    """Text extracted from one source document or document page."""

    content: str
    source_path: str
    source_name: str
    page_start: int | None = None
    page_end: int | None = None
    extraction_method: str = "text"
    heading: str | None = None


def iter_supported_files(docs_path: Path = config.SAMPLE_DOCS_PATH) -> list[Path]:
    """Return supported source files recursively in deterministic order."""
    return sorted(
        path
        for path in docs_path.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def _first_markdown_heading(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip() or None
    return None


def load_text_document(path: Path) -> list[LoadedTextBlock]:
    """Load a plain text file."""
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    return [
        LoadedTextBlock(
            content=text,
            source_path=str(path.resolve()),
            source_name=path.name,
            extraction_method="text",
        )
    ]


def load_markdown_document(path: Path) -> list[LoadedTextBlock]:
    """Load a Markdown file."""
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    return [
        LoadedTextBlock(
            content=text,
            source_path=str(path.resolve()),
            source_name=path.name,
            extraction_method="markdown",
            heading=_first_markdown_heading(text),
        )
    ]


def _needs_ocr(text: str, min_text_chars: int = config.OCR_MIN_TEXT_CHARS) -> bool:
    """Return whether extracted PDF text is too sparse and may need OCR."""
    return len(text.strip()) < min_text_chars


def _ocr_page_text(page: object) -> str:
    """Extract OCR text from a PyMuPDF page or raise a setup-focused error."""
    try:
        text_page = page.get_textpage_ocr(language=config.OCR_LANGUAGE)
        return str(page.get_text("text", textpage=text_page)).strip()
    except Exception as exc:  # pragma: no cover - depends on local Tesseract install
        raise OCRUnavailableError(
            "OCR is required for at least one PDF page, but Tesseract OCR is not available "
            "or tessdata is not configured. Install Tesseract and set TESSDATA_PREFIX on Windows."
        ) from exc


def load_pdf_document(path: Path, *, allow_ocr: bool = True) -> list[LoadedTextBlock]:
    """Load text from a PDF, using OCR only for sparse scanned pages."""
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover - dependency is installed in normal setup
        raise DocumentLoadError("PyMuPDF is required to load PDF documents.") from exc

    blocks: list[LoadedTextBlock] = []
    try:
        document = pymupdf.open(path)
    except Exception as exc:
        raise DocumentLoadError(f"Could not open PDF document: {path}") from exc

    try:
        for page_index in range(len(document)):
            page = document[page_index]
            page_number = page_index + 1
            text = str(page.get_text("text")).strip()
            method = "pdf_text"
            has_images = bool(page.get_images(full=True))

            if _needs_ocr(text) and has_images:
                if not allow_ocr:
                    raise OCRUnavailableError(
                        "OCR is required for a scanned PDF page, but OCR is disabled."
                    )
                text = _ocr_page_text(page)
                method = "ocr"

            if not text:
                continue

            blocks.append(
                LoadedTextBlock(
                    content=text,
                    source_path=str(path.resolve()),
                    source_name=path.name,
                    page_start=page_number,
                    page_end=page_number,
                    extraction_method=method,
                )
            )
    finally:
        document.close()

    return blocks


def load_docx_document(path: Path) -> list[LoadedTextBlock]:
    """Load text from a DOCX document."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is installed in normal setup
        raise DocumentLoadError("python-docx is required to load DOCX documents.") from exc

    try:
        document = Document(path)
    except Exception as exc:
        raise DocumentLoadError(f"Could not open DOCX document: {path}") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts).strip()
    if not text:
        return []

    return [
        LoadedTextBlock(
            content=text,
            source_path=str(path.resolve()),
            source_name=path.name,
            extraction_method="docx",
            heading=parts[0] if parts else None,
        )
    ]


def load_document(path: Path) -> list[LoadedTextBlock]:
    """Load one supported document by extension."""
    suffix = path.suffix.lower()
    if suffix == ".md":
        return load_markdown_document(path)
    if suffix == ".txt":
        return load_text_document(path)
    if suffix == ".pdf":
        return load_pdf_document(path)
    if suffix == ".docx":
        return load_docx_document(path)
    raise DocumentLoadError(f"Unsupported document extension: {path.suffix}")
